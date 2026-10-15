from __future__ import annotations

import csv
import io
import json
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from crosscheck.models import DocumentData, TableData, TextChunk


def _read_bytes(source: str | Path | bytes | Any, filename: str | None = None) -> tuple[bytes, str]:
    """Read raw bytes and determine filename from path, bytes, or UploadedFile."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        return path.read_bytes(), filename or path.name
    if isinstance(source, bytes):
        return source, filename or "input.dat"
    if hasattr(source, "getvalue"):
        name = filename or getattr(source, "name", "uploaded_file")
        return source.getvalue(), name
    if hasattr(source, "read"):
        name = filename or getattr(source, "name", "stream")
        data = source.read()
        if isinstance(data, str):
            data = data.encode("utf-8")
        return data, name
    raise TypeError(f"Unsupported source type: {type(source)}")


def _decode_bytes(raw: bytes) -> tuple[str, str]:
    """Decode bytes handling UTF-8 BOM (utf-8-sig) with fallback to cp1252/latin-1."""
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8-replace"


def detect_csv_delimiter(sample_text: str) -> str:
    """Detect whether a CSV uses comma, semicolon, tab, or pipe."""
    lines = [line for line in sample_text.splitlines()[:20] if line.strip()]
    if not lines:
        return ","
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines), delimiters=",;\t|")
        return dialect.delimiter
    except csv.Error:
        first_line = lines[0]
        counts = {d: first_line.count(d) for d in (",", ";", "\t", "|")}
        best = max(counts, key=counts.get)  # type: ignore[arg-type]
        return best if counts[best] > 0 else ","


def load_csv_table(raw: bytes, name: str, delimiter: str | None = None) -> TableData:
    text, encoding = _decode_bytes(raw)
    delim = delimiter or detect_csv_delimiter(text)
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    columns = [col.strip() for col in (reader.fieldnames or []) if col is not None]

    rows: list[dict[str, Any]] = []
    for raw_row in reader:
        cleaned: dict[str, Any] = {}
        for k, v in raw_row.items():
            if k is None:
                continue
            cleaned[k.strip()] = (v or "").strip()
        if any(cleaned.values()):
            rows.append(cleaned)

    return TableData(
        name=name,
        columns=columns,
        rows=rows,
        metadata={"format": "csv", "delimiter": delim, "encoding": encoding},
    )


def get_excel_sheet_names(source: str | Path | bytes | Any) -> list[str]:
    import openpyxl

    raw, _ = _read_bytes(source)
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def load_excel_table(raw: bytes, name: str, sheet_name: str | None = None) -> TableData:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        target_sheet = sheet_name if (sheet_name and sheet_name in wb.sheetnames) else wb.sheetnames[0]
        ws = wb[target_sheet]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return TableData(name=name, columns=[], rows=[], metadata={"format": "xlsx", "sheet": target_sheet})

        columns = [str(col).strip() if col is not None else f"Col_{i+1}" for i, col in enumerate(header_row)]
        rows: list[dict[str, Any]] = []
        for r in rows_iter:
            if not r or all(cell is None or str(cell).strip() == "" for cell in r):
                continue
            row_dict = {
                columns[i]: (str(r[i]).strip() if i < len(r) and r[i] is not None else "")
                for i in range(len(columns))
            }
            rows.append(row_dict)

        return TableData(
            name=name,
            columns=columns,
            rows=rows,
            metadata={"format": "xlsx", "sheet": target_sheet, "sheets": list(wb.sheetnames)},
        )
    finally:
        wb.close()


def load_json_table(raw: bytes, name: str) -> TableData:
    text, encoding = _decode_bytes(raw)
    data = json.loads(text)

    if isinstance(data, dict):
        # Find the first list value if wrapped in an object (e.g. {"users": [...]})
        list_Candidate = next((v for v in data.values() if isinstance(v, list)), None)
        if list_Candidate is not None:
            data = list_Candidate
        else:
            data = [{"Key": k, "Value": v} for k, v in data.items()]

    if not isinstance(data, list):
        data = [{"Value": str(data)}]

    columns_set: list[str] = []
    rows: list[dict[str, Any]] = []
    for item in data:
        if isinstance(item, dict):
            row: dict[str, Any] = {}
            for k, v in item.items():
                if k not in columns_set:
                    columns_set.append(k)
                row[k] = "" if v is None else (json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else str(v))
            rows.append(row)
        else:
            if "Value" not in columns_set:
                columns_set.append("Value")
            rows.append({"Value": str(item)})

    return TableData(name=name, columns=columns_set, rows=rows, metadata={"format": "json", "encoding": encoding})


def load_xml_table(raw: bytes, name: str) -> TableData:
    text, encoding = _decode_bytes(raw)
    root = ET.fromstring(text)
    columns: list[str] = []
    rows: list[dict[str, Any]] = []

    records = list(root)
    for rec in records:
        row: dict[str, Any] = {}
        for attr_k, attr_v in rec.attrib.items():
            if attr_k not in columns:
                columns.append(attr_k)
            row[attr_k] = attr_v
        for child in rec:
            tag = child.tag.split("}")[-1]
            if tag not in columns:
                columns.append(tag)
            row[tag] = (child.text or "").strip()
        if not row and rec.text and rec.text.strip():
            tag = rec.tag.split("}")[-1]
            if tag not in columns:
                columns.append(tag)
            row[tag] = rec.text.strip()
        if row:
            rows.append(row)

    return TableData(name=name, columns=columns, rows=rows, metadata={"format": "xml", "encoding": encoding})


def load_ldif_table(raw: bytes, name: str) -> TableData:
    """Parse Active Directory / LDAP LDIF exports into a structured table."""
    text, encoding = _decode_bytes(raw)
    columns: list[str] = []
    rows: list[dict[str, Any]] = []
    current: dict[str, list[str]] = {}

    def _flush(record: dict[str, list[str]]) -> None:
        if not record:
            return
        flat: dict[str, Any] = {}
        for k, vals in record.items():
            if k not in columns:
                columns.append(k)
            flat[k] = "; ".join(vals)
        rows.append(flat)

    for line in text.splitlines():
        if not line.strip():
            _flush(current)
            current = {}
            continue
        if line.startswith("#"):
            continue
        if ":" in line:
            key, val = line.split(":", 1)
            key = key.strip()
            val = val.lstrip(": ").strip()
            current.setdefault(key, []).append(val)

    _flush(current)
    return TableData(name=name, columns=columns, rows=rows, metadata={"format": "ldif", "encoding": encoding})


def load_lines_table(raw: bytes, name: str) -> TableData:
    """Load a plain text file with one item per line as a 1-column table."""
    text, encoding = _decode_bytes(raw)
    rows = [{"Value": line.strip()} for line in text.splitlines() if line.strip()]
    return TableData(name=name, columns=["Value"], rows=rows, metadata={"format": "txt", "encoding": encoding})


def load_table(
    source: str | Path | bytes | Any,
    filename: str | None = None,
    delimiter: str | None = None,
    sheet_name: str | None = None,
) -> TableData:
    """Universal table loader supporting CSV, TSV, XLSX, JSON, XML, LDIF, and TXT lists."""
    raw, name = _read_bytes(source, filename)
    ext = Path(name).suffix.lower()

    if ext in (".xlsx", ".xlsm"):
        return load_excel_table(raw, name, sheet_name=sheet_name)
    if ext == ".json":
        return load_json_table(raw, name)
    if ext == ".xml":
        return load_xml_table(raw, name)
    if ext == ".ldif":
        return load_ldif_table(raw, name)
    if ext == ".tsv":
        return load_csv_table(raw, name, delimiter=delimiter or "\t")
    if ext == ".csv":
        return load_csv_table(raw, name, delimiter=delimiter)

    # For .txt or unknown extensions, check if it looks like a delimited table with headers
    text, _ = _decode_bytes(raw)
    first_line = next((l for l in text.splitlines() if l.strip()), "")
    if delimiter or any(d in first_line for d in (";", ",", "\t")):
        return load_csv_table(raw, name, delimiter=delimiter)
    return load_lines_table(raw, name)


def load_document(source: str | Path | bytes | Any, filename: str | None = None) -> DocumentData:
    """Universal document loader supporting PDF, DOCX, TXT, LOG, CSV, JSON, XML."""
    raw, name = _read_bytes(source, filename)
    ext = Path(name).suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        chunks: list[TextChunk] = []
        for idx, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                chunks.append(TextChunk(location=f"Sida {idx}", text=page_text))
        return DocumentData(name=name, chunks=chunks, metadata={"format": "pdf", "pages": len(reader.pages)})

    if ext == ".docx":
        import docx

        doc = docx.Document(io.BytesIO(raw))
        chunks = []
        for idx, para in enumerate(doc.paragraphs, start=1):
            t = para.text.strip()
            if t:
                chunks.append(TextChunk(location=f"Stycke {idx}", text=t))
        # Also include text inside tables in the Word document
        for t_idx, table in enumerate(doc.tables, start=1):
            for r_idx, row in enumerate(table.rows, start=1):
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    chunks.append(TextChunk(location=f"Tabell {t_idx} Rad {r_idx}", text=row_text))
        return DocumentData(name=name, chunks=chunks, metadata={"format": "docx"})

    if ext in (".xlsx", ".xlsm"):
        table_data = load_excel_table(raw, name)
        chunks = []
        for idx, row in enumerate(table_data.rows, start=2):
            line_str = " | ".join(f"{k}: {v}" for k, v in row.items() if v)
            if line_str:
                chunks.append(TextChunk(location=f"Rad {idx}", text=line_str))
        return DocumentData(name=name, chunks=chunks, metadata={"format": "xlsx"})

    # Default: line-based text document (.txt, .log, .csv, .md, .json, .xml, etc.)
    text, encoding = _decode_bytes(raw)
    chunks = [
        TextChunk(location=f"Rad {idx}", text=line)
        for idx, line in enumerate(text.splitlines(), start=1)
        if line.strip()
    ]
    return DocumentData(name=name, chunks=chunks, metadata={"format": ext.lstrip(".") or "txt", "encoding": encoding})
