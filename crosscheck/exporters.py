from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from crosscheck.models import TableMatchResult, TextSearchResult


def rows_to_csv_bytes(rows: list[dict[str, Any]], columns: list[str], delimiter: str = ";") -> bytes:
    """Convert a list of row dicts to UTF-8-SIG (Excel-friendly) CSV bytes."""
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=columns, delimiter=delimiter, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({col: row.get(col, "") for col in columns})
    return buf.getvalue().encode("utf-8-sig")


def match_result_to_excel_bytes(result: TableMatchResult) -> bytes:
    """Export a TableMatchResult into a multi-sheet Excel (.xlsx) workbook."""
    import openpyxl

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)  # type: ignore[arg-type]

    def _add_sheet(title: str, columns: list[str], rows: list[dict[str, Any]]) -> None:
        ws = wb.create_sheet(title=title)
        if not columns:
            return
        ws.append(columns)
        for row in rows:
            ws.append([str(row.get(c, "") or "") for c in columns])

    _add_sheet("Finns i båda (Match)", result.matched_columns, result.matched)
    _add_sheet("Saknas i B (Bara i A)", result.columns_a, result.only_in_a)
    _add_sheet("Saknas i A (Bara i B)", result.columns_b, result.only_in_b)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def text_search_to_excel_bytes(result: TextSearchResult) -> bytes:
    """Export a TextSearchResult into a multi-sheet Excel (.xlsx) workbook."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    ws_hits = wb.create_sheet(title="Träffar i dokument")
    ws_hits.append(["Sökord", "Plats", "Textutdrag"])
    for hit in result.hits:
        ws_hits.append([hit.term, hit.location, hit.snippet])

    ws_found = wb.create_sheet(title="Hittade sökord")
    ws_found.append(["Sökord", "Status"])
    for t in result.found_terms:
        ws_found.append([t, "Hittad"])

    ws_missing = wb.create_sheet(title="Saknas i dokument")
    ws_missing.append(["Sökord", "Status"])
    for t in result.missing_terms:
        ws_missing.append([t, "Saknas"])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def save_bytes(data: bytes, path: str | Path) -> Path:
    out_path = Path(path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_bytes(data)
    return out_path
